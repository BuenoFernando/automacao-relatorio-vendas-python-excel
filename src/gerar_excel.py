from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

# ==============================
# CONFIGURAÇÕES
# ==============================

# Caminho da pasta do projeto
BASE_DIR = Path(__file__).resolve().parent.parent

# Arquivos de entrada e saída
arquivo_entrada = BASE_DIR / "dados" / "vendas.xlsx"
arquivo_saida = BASE_DIR / "relatorios" / "relatorio_vendas.xlsx"

# ==============================
# LEITURA DA BASE
# ==============================

df = pd.read_excel(arquivo_entrada)

# Apenas vendas concluídas
df = df[df["Status"] == "Concluído"].copy()

# ==============================
# INDICADORES
# ==============================

faturamento_total = df["Faturamento"].sum()
quantidade_vendida = df["Quantidade"].sum()
total_pedidos = len(df)
ticket_medio = faturamento_total / total_pedidos

# ==============================
# TABELAS DE ANÁLISE
# ==============================

por_produto = (
    df.groupby("Produto")
    .agg(
        Quantidade=("Quantidade", "sum"),
        Faturamento=("Faturamento", "sum")
    )
    .sort_values("Faturamento", ascending=False)
    .reset_index()
)

por_regiao = (
    df.groupby("Região")
    .agg(
        Quantidade=("Quantidade", "sum"),
        Faturamento=("Faturamento", "sum")
    )
    .sort_values("Faturamento", ascending=False)
    .reset_index()
)

por_vendedor = (
    df.groupby("Vendedor")
    .agg(
        Quantidade=("Quantidade", "sum"),
        Faturamento=("Faturamento", "sum")
    )
    .sort_values("Faturamento", ascending=False)
    .reset_index()
)

por_pagamento = (
    df.groupby("Forma de Pagamento")
    .agg(
        Quantidade=("Quantidade", "sum"),
        Faturamento=("Faturamento", "sum")
    )
    .sort_values("Faturamento", ascending=False)
    .reset_index()
)

# ==============================
# RESUMO
# ==============================

resumo = pd.DataFrame({
    "Indicador": [
        "Faturamento Total",
        "Quantidade Vendida",
        "Pedidos Concluídos",
        "Ticket Médio"
    ],
    "Valor": [
        faturamento_total,
        quantidade_vendida,
        total_pedidos,
        ticket_medio
    ]
})

# ==============================
# CRIAR EXCEL
# ==============================

with pd.ExcelWriter(arquivo_saida, engine="openpyxl") as writer:

    resumo.to_excel(
        writer,
        sheet_name="Resumo",
        index=False
    )

    por_produto.to_excel(
        writer,
        sheet_name="Por Produto",
        index=False
    )

    por_regiao.to_excel(
        writer,
        sheet_name="Por Região",
        index=False
    )

    por_vendedor.to_excel(
        writer,
        sheet_name="Por Vendedor",
        index=False
    )

    por_pagamento.to_excel(
        writer,
        sheet_name="Por Pagamento",
        index=False
    )

# ==============================
# FORMATAÇÃO
# ==============================

wb = load_workbook(arquivo_saida)

for ws in wb.worksheets:

    # Cabeçalho
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Largura das colunas
    for column in ws.columns:

        maior = 0
        coluna = get_column_letter(column[0].column)

        for cell in column:

            if cell.value is not None:

                tamanho = len(str(cell.value))

                if tamanho > maior:
                    maior = tamanho

        ws.column_dimensions[coluna].width = maior + 3

    # Congelar cabeçalho
    ws.freeze_panes = "A2"


# ==============================
# FORMATAÇÃO DE MOEDA
# ==============================

ws = wb["Resumo"]

ws["B2"].number_format = 'R$ #,##0.00'
ws["B5"].number_format = 'R$ #,##0.00'


for nome_planilha in [
    "Por Produto",
    "Por Região",
    "Por Vendedor",
    "Por Pagamento"
]:

    ws = wb[nome_planilha]

    for cell in ws["C"][1:]:
        cell.number_format = 'R$ #,##0.00'


# ==============================
# GRÁFICO — PRODUTOS
# ==============================

ws = wb["Por Produto"]

grafico_produto = BarChart()

grafico_produto.title = "Faturamento por Produto"
grafico_produto.y_axis.title = "Faturamento"
grafico_produto.x_axis.title = "Produto"

dados = Reference(
    ws,
    min_col=3,
    min_row=1,
    max_row=ws.max_row
)

categorias = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=ws.max_row
)

grafico_produto.add_data(
    dados,
    titles_from_data=True
)

grafico_produto.set_categories(categorias)

grafico_produto.height = 8
grafico_produto.width = 14

ws.add_chart(
    grafico_produto,
    "E2"
)


# ==============================
# GRÁFICO — REGIÕES
# ==============================

ws = wb["Por Região"]

grafico_regiao = BarChart()

grafico_regiao.title = "Faturamento por Região"
grafico_regiao.y_axis.title = "Faturamento"
grafico_regiao.x_axis.title = "Região"

dados = Reference(
    ws,
    min_col=3,
    min_row=1,
    max_row=ws.max_row
)

categorias = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=ws.max_row
)

grafico_regiao.add_data(
    dados,
    titles_from_data=True
)

grafico_regiao.set_categories(categorias)

grafico_regiao.height = 8
grafico_regiao.width = 14

ws.add_chart(
    grafico_regiao,
    "E2"
)


# ==============================
# GRÁFICO — VENDEDORES
# ==============================

ws = wb["Por Vendedor"]

grafico_vendedor = BarChart()

grafico_vendedor.title = "Faturamento por Vendedor"
grafico_vendedor.y_axis.title = "Faturamento"
grafico_vendedor.x_axis.title = "Vendedor"

dados = Reference(
    ws,
    min_col=3,
    min_row=1,
    max_row=ws.max_row
)

categorias = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=ws.max_row
)

grafico_vendedor.add_data(
    dados,
    titles_from_data=True
)

grafico_vendedor.set_categories(categorias)

grafico_vendedor.height = 8
grafico_vendedor.width = 14

ws.add_chart(
    grafico_vendedor,
    "E2"
)

# ==========================================
# DASHBOARD
# ==========================================

from openpyxl.styles import Font, PatternFill, Border, Side

# Criar aba Dashboard
if "Dashboard" in wb.sheetnames:
    del wb["Dashboard"]

dashboard = wb.create_sheet("Dashboard", 0)

# Título
dashboard["B2"] = "DASHBOARD DE VENDAS"
dashboard["B2"].font = Font(size=20, bold=True)
dashboard["B2"].alignment = Alignment(horizontal="center")

dashboard.merge_cells("B2:J2")

# KPIs
dashboard["B4"] = "Faturamento Total"
dashboard["D4"] = "Quantidade Vendida"
dashboard["F4"] = "Pedidos Concluídos"
dashboard["H4"] = "Ticket Médio"

dashboard["B5"] = faturamento_total
dashboard["D5"] = quantidade_vendida
dashboard["F5"] = total_pedidos
dashboard["H5"] = ticket_medio

# Mesclar células dos cartões
dashboard.merge_cells("B4:C4")
dashboard.merge_cells("B5:C5")

dashboard.merge_cells("D4:E4")
dashboard.merge_cells("D5:E5")

dashboard.merge_cells("F4:G4")
dashboard.merge_cells("F5:G5")

dashboard.merge_cells("H4:I4")
dashboard.merge_cells("H5:I5")

# ==========================================
# FORMATAÇÃO PROFISSIONAL DOS KPIs
# ==========================================

# Estilo dos cartões
preenchimento = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

borda = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)

# Títulos dos cartões
for celula in ["B4", "D4", "F4", "H4"]:
    dashboard[celula].font = Font(
        size=11,
        bold=True
    )

    dashboard[celula].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    dashboard[celula].fill = preenchimento
    dashboard[celula].border = borda

# Valores dos cartões
for celula in ["B5", "D5", "F5", "H5"]:
    dashboard[celula].font = Font(
        size=16,
        bold=True
    )

    dashboard[celula].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    dashboard[celula].fill = preenchimento
    dashboard[celula].border = borda

# Formatação monetária
dashboard["B5"].number_format = 'R$ #,##0.00'
dashboard["H5"].number_format = 'R$ #,##0.00'

# Bordas e preenchimento nas áreas dos cartões
intervalos_cartoes = [
    "B4:C5",
    "D4:E5",
    "F4:G5",
    "H4:I5"
]

for intervalo in intervalos_cartoes:
    for linha in dashboard[intervalo]:
        for celula in linha:
            celula.fill = preenchimento
            celula.border = borda

# Altura dos cartões
dashboard.row_dimensions[4].height = 28
dashboard.row_dimensions[5].height = 35

# Ajuste das larguras
for coluna in ["B", "C", "D", "E", "F", "G", "H", "I", "J"]:
    dashboard.column_dimensions[coluna].width = 16

# Altura das linhas dos KPIs
dashboard.row_dimensions[4].height = 25
dashboard.row_dimensions[5].height = 30

# ==========================================
# GRÁFICOS DO DASHBOARD
# ==========================================

# Gráfico de faturamento por produto
grafico_produto_dashboard = BarChart()
grafico_produto_dashboard.title = "Faturamento por Produto"
grafico_produto_dashboard.y_axis.title = "Faturamento"
grafico_produto_dashboard.x_axis.title = "Produto"

ws_produto = wb["Por Produto"]

dados = Reference(
    ws_produto,
    min_col=3,
    min_row=1,
    max_row=ws_produto.max_row
)

categorias = Reference(
    ws_produto,
    min_col=1,
    min_row=2,
    max_row=ws_produto.max_row
)

grafico_produto_dashboard.add_data(
    dados,
    titles_from_data=True
)

grafico_produto_dashboard.set_categories(categorias)

grafico_produto_dashboard.height = 8
grafico_produto_dashboard.width = 14

dashboard.add_chart(
    grafico_produto_dashboard,
    "B8"
)


# Gráfico de faturamento por região
grafico_regiao_dashboard = BarChart()
grafico_regiao_dashboard.title = "Faturamento por Região"
grafico_regiao_dashboard.y_axis.title = "Faturamento"
grafico_regiao_dashboard.x_axis.title = "Região"

ws_regiao = wb["Por Região"]

dados = Reference(
    ws_regiao,
    min_col=3,
    min_row=1,
    max_row=ws_regiao.max_row
)

categorias = Reference(
    ws_regiao,
    min_col=1,
    min_row=2,
    max_row=ws_regiao.max_row
)

grafico_regiao_dashboard.add_data(
    dados,
    titles_from_data=True
)

grafico_regiao_dashboard.set_categories(categorias)

grafico_regiao_dashboard.height = 8
grafico_regiao_dashboard.width = 14

dashboard.add_chart(
    grafico_regiao_dashboard,
    "J8"
)


# Gráfico de faturamento por vendedor
grafico_vendedor_dashboard = BarChart()
grafico_vendedor_dashboard.title = "Faturamento por Vendedor"
grafico_vendedor_dashboard.y_axis.title = "Faturamento"
grafico_vendedor_dashboard.x_axis.title = "Vendedor"

ws_vendedor = wb["Por Vendedor"]

dados = Reference(
    ws_vendedor,
    min_col=3,
    min_row=1,
    max_row=ws_vendedor.max_row
)

categorias = Reference(
    ws_vendedor,
    min_col=1,
    min_row=2,
    max_row=ws_vendedor.max_row
)

grafico_vendedor_dashboard.add_data(
    dados,
    titles_from_data=True
)

grafico_vendedor_dashboard.set_categories(categorias)

grafico_vendedor_dashboard.height = 8
grafico_vendedor_dashboard.width = 14

dashboard.add_chart(
    grafico_vendedor_dashboard,
    "F25"
)

wb.save(arquivo_saida)

print("====================================")
print("RELATÓRIO EXCEL GERADO COM SUCESSO!")
print("====================================")
print(f"Arquivo: {arquivo_saida}")
print("Gráficos adicionados com sucesso!")